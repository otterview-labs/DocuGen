#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
支持多种文档格式的深度解析
"""

import os
import re
import logging
from typing import Dict, List, Any, Optional, Union, Tuple
from pathlib import Path
from io import BytesIO
import json

# 文档处理库
import pdfplumber
from docx import Document
import pandas as pd
from openpyxl import load_workbook
from PIL import Image
import markdown
from bs4 import BeautifulSoup

# 文本处理
import jieba
import jieba.posseg as pseg
from collections import Counter

from .file_parse import Parser
from .file_quckly_parse import FileToMarkdown

logger = logging.getLogger(__name__)


class EnhancedDocumentParser:
    """增强文档解析器 - 参考深度设计"""

    def __init__(self):
        self.file_parser = Parser()
        self.quick_parser = FileToMarkdown()

        # 支持的文档类型
        self.supported_formats = {
            '.pdf': self._parse_pdf,
            '.docx': self._parse_docx,
            '.doc': self._parse_doc,
            '.xlsx': self._parse_excel,
            '.xls': self._parse_excel,
            '.csv': self._parse_csv,
            '.txt': self._parse_txt,
            '.md': self._parse_markdown,
            '.html': self._parse_html,
            '.json': self._parse_json,
            '.jpg': self._parse_image,
            '.jpeg': self._parse_image,
            '.png': self._parse_image,
            '.bmp': self._parse_image,
            '.tiff': self._parse_image,
        }

    def parse_document(self, file_path: str, **kwargs) -> Dict[str, Any]:
        """
        解析文档主入口

        Args:
            file_path: 文件路径
            **kwargs: 额外参数

        Returns:
            解析结果字典
        """
        try:
            file_path = Path(file_path)
            if not file_path.exists():
                return {"error": f"文件不存在: {file_path}"}

            file_ext = file_path.suffix.lower()

            if file_ext not in self.supported_formats:
                return {"error": f"不支持的文件格式: {file_ext}"}

            # 调用对应的解析器
            parser_func = self.supported_formats[file_ext]
            result = parser_func(file_path, **kwargs)

            # 添加元数据
            result["metadata"] = {
                "file_path": str(file_path),
                "file_name": file_path.name,
                "file_size": file_path.stat().st_size,
                "file_extension": file_ext,
                "parser_type": "enhanced_parser"
            }

            return result

        except Exception as e:
            logger.error(f"解析文档失败: {e}")
            return {"error": str(e)}

    def _parse_pdf(self, file_path: Path, **kwargs) -> Dict[str, Any]:
        """PDF文档解析 - 参考深度的PDF解析器"""
        try:
            result = {
                "content": "",
                "tables": [],
                "images": [],
                "structure": {
                    "pages": [],
                    "sections": [],
                    "headers": [],
                    "footers": []
                },
                "metadata": {}
            }

            # 使用pdfplumber解析
            with pdfplumber.open(file_path) as pdf:
                result["metadata"]["total_pages"] = len(pdf.pages)

                for page_num, page in enumerate(pdf.pages):
                    page_content = {
                        "page_number": page_num + 1,
                        "text": "",
                        "tables": [],
                        "images": [],
                        "layout": {}
                    }

                    # 提取文本
                    page_text = page.extract_text()
                    if page_text:
                        page_content["text"] = page_text
                        result["content"] += f"\n\n--- 第{page_num + 1}页 ---\n\n{page_text}"

                    # 提取表格
                    tables = page.extract_tables()
                    for table_num, table in enumerate(tables):
                        if table and any(any(cell for cell in row) for row in table):
                            table_data = {
                                "page": page_num + 1,
                                "table_index": table_num,
                                "data": table,
                                "html": self._table_to_html(table)
                            }
                            page_content["tables"].append(table_data)
                            result["tables"].append(table_data)

                    # 提取图像信息
                    if page.images:
                        for img_num, img in enumerate(page.images):
                            img_data = {
                                "page": page_num + 1,
                                "image_index": img_num,
                                "bbox": img["bbox"],
                                "width": img["width"],
                                "height": img["height"]
                            }
                            page_content["images"].append(img_data)
                            result["images"].append(img_data)

                    result["structure"]["pages"].append(page_content)

            # 智能分段
            result["segments"] = self._segment_content(result["content"])

            return result

        except Exception as e:
            logger.error(f"PDF解析失败: {e}")
            return {"error": f"PDF解析失败: {str(e)}"}

    def _parse_docx(self, file_path: Path, **kwargs) -> Dict[str, Any]:
        """Word文档解析 - 参考深度的DOCX解析器"""
        try:
            result = {
                "content": "",
                "tables": [],
                "structure": {
                    "paragraphs": [],
                    "sections": [],
                    "headers": [],
                    "footers": []
                },
                "metadata": {}
            }

            doc = Document(file_path)

            # 提取段落
            for para in doc.paragraphs:
                if para.text.strip():
                    result["content"] += para.text + "\n"
                    result["structure"]["paragraphs"].append({
                        "text": para.text,
                        "style": para.style.name,
                        "alignment": str(para.alignment)
                    })

            # 提取表格
            for table_num, table in enumerate(doc.tables):
                table_data = self._extract_docx_table(table)
                result["tables"].append(table_data)

            # 智能分段
            result["segments"] = self._segment_content(result["content"])

            return result

        except Exception as e:
            logger.error(f"DOCX解析失败: {e}")
            return {"error": f"DOCX解析失败: {str(e)}"}

    def _parse_excel(self, file_path: Path, **kwargs) -> Dict[str, Any]:
        """Excel文档解析 - 参考深度的Excel解析器"""
        try:
            result = {
                "content": "",
                "sheets": [],
                "tables": [],
                "metadata": {}
            }

            # 尝试使用openpyxl
            try:
                wb = load_workbook(file_path, data_only=True)
                result["metadata"]["sheets"] = wb.sheetnames

                for sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]
                    sheet_data = {
                        "name": sheet_name,
                        "rows": [],
                        "html": self._worksheet_to_html(ws, sheet_name)
                    }

                    # 提取数据
                    for row in ws.iter_rows(values_only=True):
                        if any(cell is not None for cell in row):
                            sheet_data["rows"].append(row)

                    result["sheets"].append(sheet_data)
                    result["content"] += f"\n\n--- {sheet_name} ---\n\n"
                    result["content"] += self._worksheet_to_text(ws)

            except Exception as e:
                # 回退到pandas
                logger.warning(f"openpyxl解析失败，使用pandas: {e}")
                df = pd.read_excel(file_path, sheet_name=None)

                for sheet_name, sheet_df in df.items():
                    sheet_data = {
                        "name": sheet_name,
                        "data": sheet_df.to_dict('records'),
                        "html": sheet_df.to_html()
                    }
                    result["sheets"].append(sheet_data)
                    result["content"] += f"\n\n--- {sheet_name} ---\n\n"
                    result["content"] += sheet_df.to_string()

            return result

        except Exception as e:
            logger.error(f"Excel解析失败: {e}")
            return {"error": f"Excel解析失败: {str(e)}"}

    def _parse_csv(self, file_path: Path, **kwargs) -> Dict[str, Any]:
        """CSV文件解析"""
        try:
            result = {
                "content": "",
                "data": [],
                "metadata": {}
            }

            df = pd.read_csv(file_path)
            result["data"] = df.to_dict('records')
            result["content"] = df.to_string()
            result["metadata"]["columns"] = list(df.columns)
            result["metadata"]["rows"] = len(df)

            return result

        except Exception as e:
            logger.error(f"CSV解析失败: {e}")
            return {"error": f"CSV解析失败: {str(e)}"}

    def _parse_txt(self, file_path: Path, **kwargs) -> Dict[str, Any]:
        """文本文件解析"""
        try:
            result = {
                "content": "",
                "segments": [],
                "metadata": {}
            }

            # 尝试不同编码
            encodings = ['utf-8', 'gbk', 'gb2312', 'latin-1']
            content = None

            for encoding in encodings:
                try:
                    with open(file_path, 'r', encoding=encoding) as f:
                        content = f.read()
                    break
                except UnicodeDecodeError:
                    continue

            if content is None:
                return {"error": "无法解码文本文件"}

            result["content"] = content
            result["segments"] = self._segment_content(content)
            result["metadata"]["encoding"] = encoding

            return result

        except Exception as e:
            logger.error(f"TXT解析失败: {e}")
            return {"error": f"TXT解析失败: {str(e)}"}

    def _parse_markdown(self, file_path: Path, **kwargs) -> Dict[str, Any]:
        """Markdown文件解析"""
        try:
            result = {
                "content": "",
                "html": "",
                "structure": {
                    "headers": [],
                    "links": [],
                    "images": []
                },
                "metadata": {}
            }

            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()

            result["content"] = content
            result["html"] = markdown.markdown(content)

            # 提取结构信息
            lines = content.split('\n')
            for line in lines:
                if line.startswith('#'):
                    result["structure"]["headers"].append(line)
                elif '![' in line and '](' in line:
                    result["structure"]["images"].append(line)
                elif '[' in line and '](' in line:
                    result["structure"]["links"].append(line)

            result["segments"] = self._segment_content(content)

            return result

        except Exception as e:
            logger.error(f"Markdown解析失败: {e}")
            return {"error": f"Markdown解析失败: {str(e)}"}

    def _parse_html(self, file_path: Path, **kwargs) -> Dict[str, Any]:
        """HTML文件解析"""
        try:
            result = {
                "content": "",
                "structure": {
                    "titles": [],
                    "links": [],
                    "images": []
                },
                "metadata": {}
            }

            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()

            soup = BeautifulSoup(content, 'html.parser')

            # 提取文本
            result["content"] = soup.get_text()

            # 提取结构信息
            for title in soup.find_all(['h1', 'h2', 'h3', 'h4', 'h5', 'h6']):
                result["structure"]["titles"].append(title.get_text())

            for link in soup.find_all('a'):
                result["structure"]["links"].append({
                    "text": link.get_text(),
                    "href": link.get('href')
                })

            for img in soup.find_all('img'):
                result["structure"]["images"].append({
                    "alt": img.get('alt'),
                    "src": img.get('src')
                })

            result["segments"] = self._segment_content(result["content"])

            return result

        except Exception as e:
            logger.error(f"HTML解析失败: {e}")
            return {"error": f"HTML解析失败: {str(e)}"}

    def _parse_json(self, file_path: Path, **kwargs) -> Dict[str, Any]:
        """JSON文件解析"""
        try:
            result = {
                "content": "",
                "data": {},
                "metadata": {}
            }

            with open(file_path, 'r', encoding='utf-8') as f:
                data = json.load(f)

            result["data"] = data
            result["content"] = json.dumps(data, ensure_ascii=False, indent=2)
            result["metadata"]["type"] = type(data).__name__

            return result

        except Exception as e:
            logger.error(f"JSON解析失败: {e}")
            return {"error": f"JSON解析失败: {str(e)}"}

    def _parse_image(self, file_path: Path, **kwargs) -> Dict[str, Any]:
        """图像文件解析"""
        try:
            result = {
                "content": "",
                "metadata": {}
            }

            with Image.open(file_path) as img:
                result["metadata"] = {
                    "format": img.format,
                    "mode": img.mode,
                    "size": img.size,
                    "width": img.width,
                    "height": img.height
                }

                # 如果有EXIF信息
                if hasattr(img, '_getexif') and img._getexif():
                    result["metadata"]["exif"] = img._getexif()

            # 注意：图像内容解析需要OCR服务
            # 这里可以集成OCR功能
            result["content"] = f"图像文件: {file_path.name}"

            return result

        except Exception as e:
            logger.error(f"图像解析失败: {e}")
            return {"error": f"图像解析失败: {str(e)}"}

    def _parse_doc(self, file_path: Path, **kwargs) -> Dict[str, Any]:
        """DOC文件解析 - 使用现有解析器"""
        try:
            return self.quick_parser.parse_file(str(file_path))
        except Exception as e:
            logger.error(f"DOC解析失败: {e}")
            return {"error": f"DOC解析失败: {str(e)}"}

    def _segment_content(self, content: str) -> List[Dict[str, Any]]:
        """智能内容分段 - 参考深度的分段策略"""
        segments = []

        if not content:
            return segments

        # 按段落分割
        paragraphs = content.split('\n\n')

        for i, para in enumerate(paragraphs):
            para = para.strip()
            if not para:
                continue

            # 分析段落类型
            segment_type = self._analyze_segment_type(para)

            segments.append({
                "id": i,
                "content": para,
                "type": segment_type,
                "length": len(para),
                "tokens": len(para.split())  # 简单token估算
            })

        return segments

    def _analyze_segment_type(self, text: str) -> str:
        """分析文本段类型"""
        if not text:
            return "empty"

        # 标题检测
        if text.startswith('#') or (len(text) < 100 and text.isupper()):
            return "header"

        # 列表检测
        if re.match(r'^[\s]*[•\-\*\d]+\.?\s', text):
            return "list"

        # 表格检测
        if '\t' in text or '|' in text:
            return "table"

        # 代码检测
        if text.startswith('```') or text.startswith('    '):
            return "code"

        # 长文本
        if len(text) > 500:
            return "paragraph"

        return "text"

    def _table_to_html(self, table_data: List[List[str]]) -> str:
        """将表格数据转换为HTML"""
        if not table_data:
            return ""

        html = "<table border='1'>\n"

        for i, row in enumerate(table_data):
            html += "<tr>\n"
            tag = "th" if i == 0 else "td"

            for cell in row:
                cell_text = str(cell) if cell is not None else ""
                html += f"<{tag}>{cell_text}</{tag}>\n"

            html += "</tr>\n"

        html += "</table>"
        return html

    def _extract_docx_table(self, table) -> Dict[str, Any]:
        """提取Word表格数据"""
        table_data = {
            "data": [],
            "html": "",
            "metadata": {}
        }

        # 提取表格数据
        for row in table.rows:
            row_data = []
            for cell in row.cells:
                row_data.append(cell.text)
            table_data["data"].append(row_data)

        # 生成HTML
        if table_data["data"]:
            table_data["html"] = self._table_to_html(table_data["data"])

        return table_data

    def _worksheet_to_html(self, ws, sheet_name: str) -> str:
        """将工作表转换为HTML"""
        html = f"<table border='1'><caption>{sheet_name}</caption>\n"

        for row in ws.iter_rows():
            html += "<tr>\n"
            for cell in row:
                value = cell.value if cell.value is not None else ""
                html += f"<td>{value}</td>\n"
            html += "</tr>\n"

        html += "</table>"
        return html

    def _worksheet_to_text(self, ws) -> str:
        """将工作表转换为文本"""
        text = ""
        for row in ws.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                text += "\t".join(str(cell) if cell is not None else "" for cell in row) + "\n"
        return text

    def get_supported_formats(self) -> List[str]:
        """获取支持的文档格式"""
        return list(self.supported_formats.keys())

    def is_supported(self, file_path: str) -> bool:
        """检查文件格式是否支持"""
        file_ext = Path(file_path).suffix.lower()
        return file_ext in self.supported_formats
