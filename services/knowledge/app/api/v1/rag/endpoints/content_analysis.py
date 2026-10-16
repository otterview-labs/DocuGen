#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
内容分析API接口 - 摘要和标签提取
"""

import json
import os
import logging
from typing import Dict, List, Optional, Union
from datetime import datetime
from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File, Form
from sqlalchemy.orm import Session
from pydantic import BaseModel, ConfigDict, Field

from app.db.session import get_db
from app.db.models.rag import RAGContent
from app.db.models import User
from app.services.content_analysis_service import get_content_analysis_service
from app.services.enhanced_rag_service import get_enhanced_rag_service

logger = logging.getLogger(__name__)

router = APIRouter()


async def _extract_text_from_file(file_path: str, file_name: str) -> str:
    """从不同格式的文件中提取文本内容"""
    import os
    from pathlib import Path

    file_ext = Path(file_name).suffix.lower()

    try:
        if file_ext == '.txt':
            # 纯文本文件
            encodings = ['utf-8', 'gbk', 'gb2312', 'utf-16']
            for encoding in encodings:
                try:
                    with open(file_path, 'r', encoding=encoding) as f:
                        return f.read()
                except UnicodeDecodeError:
                    continue
            raise Exception("无法识别文件编码")

        elif file_ext in ['.docx', '.doc']:
            # Word文档
            try:
                from docx import Document
                doc = Document(file_path)
                text_content = ""
                for paragraph in doc.paragraphs:
                    text_content += paragraph.text + "\n"
                return text_content
            except ImportError:
                # 如果没有python-docx库，返回基础信息
                return f"Word文档: {file_name} (需要安装python-docx库来提取完整内容)"
            except Exception as e:
                return f"Word文档读取失败: {str(e)}"

        elif file_ext == '.pdf':
            # PDF文件
            try:
                import PyPDF2
                with open(file_path, 'rb') as f:
                    reader = PyPDF2.PdfReader(f)
                    text_content = ""
                    for page in reader.pages:
                        text_content += page.extract_text() + "\n"
                return text_content
            except ImportError:
                return f"PDF文档: {file_name} (需要安装PyPDF2库来提取完整内容)"
            except Exception as e:
                return f"PDF文档读取失败: {str(e)}"

        elif file_ext in ['.pptx', '.ppt']:
            # PowerPoint文档
            try:
                from pptx import Presentation
                prs = Presentation(file_path)
                text_content = ""
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text"):
                            text_content += shape.text + "\n"
                return text_content
            except ImportError:
                return f"PowerPoint文档: {file_name} (需要安装python-pptx库来提取完整内容)"
            except Exception as e:
                return f"PowerPoint文档读取失败: {str(e)}"

        else:
            # 其他文件类型，尝试作为文本读取
            encodings = ['utf-8', 'gbk', 'gb2312']
            for encoding in encodings:
                try:
                    with open(file_path, 'r', encoding=encoding) as f:
                        return f.read()
                except UnicodeDecodeError:
                    continue
                except Exception:
                    break

            # 如果都失败了，返回文件基础信息
            return f"文档: {file_name} (文件类型: {file_ext})"

    except Exception as e:
        logger.error(f"文件读取失败 {file_path}: {e}")
        return f"文档: {file_name} (读取时发生错误: {str(e)})"


class ContentAnalysisRequest(BaseModel):
    """内容分析请求模型（支持 resource_id，兼容 content_id）"""
    resource_id: Optional[Union[str, List[str]]] = Field(
        None, description="资源ID，推荐使用；可为单个或列表"
    )
    content_id: Optional[Union[str, List[str]]] = Field(
        None, description="内容ID（兼容旧版）；可为单个或列表"
    )
    analysis_type: str = Field("both", description="分析类型: summary, tags, both")

    model_config = ConfigDict(
        json_schema_extra={
            "examples": [
                {
                    "resource_id": "res_123",
                    "analysis_type": "both"
                },
                {
                    "resource_id": ["res_123", "res_456"],
                    "analysis_type": "summary"
                },
                {
                    "content_id": "content_123",  # 兼容旧用法
                    "analysis_type": "both"
                }
            ]
        }
    )


class ContentAnalysisResponse(BaseModel):
    """内容分析响应模型"""
    content_id: str
    summary: Optional[str] = None
    tags: Optional[List[str]] = None
    analysis_type: str


def get_content_analysis_service_instance():
    """获取内容分析服务实例"""
    return get_content_analysis_service()


@router.post("/analyze", response_model=Dict)
async def analyze_content(
    request: ContentAnalysisRequest,
    db: Session = Depends(get_db)
):
    """
    分析指定内容的摘要和标签（支持单个或批量分析）

    支持的分析类型：
    - summary: 仅生成摘要
    - tags: 仅提取标签
    - both: 同时生成摘要和标签

    支持的输入格式：
    - 单个内容: {"content_id": "content_123", "analysis_type": "both"}
    - 批量内容: {"content_id": ["content_123", "content_456"], "analysis_type": "both"}
    """
    try:
        # 优先使用 resource_id（新用法）；否则回退 content_id（旧用法）
        ids = request.resource_id if request.resource_id is not None else request.content_id
        if ids is None:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="必须提供 resource_id 或 content_id")

        if isinstance(ids, str):
            # 单个
            if request.resource_id is not None:
                return await _analyze_single_resource(ids, request.analysis_type, db)
            else:
                return await _analyze_single_content(ids, request.analysis_type, db)
        elif isinstance(ids, list):
            # 批量
            if request.resource_id is not None:
                return await _analyze_batch_resource(ids, request.analysis_type, db)
            else:
                return await _analyze_batch_content(ids, request.analysis_type, db)
        else:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="ID 参数必须是字符串或字符串列表")

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"内容分析失败: {str(e)}"
        )


async def _analyze_single_content(content_id: str, analysis_type: str, db: Session) -> Dict:
    """分析单个内容 - 使用增强的共享分析服务"""
    try:
        # 查找内容记录
        content = db.query(RAGContent).filter(RAGContent.id == content_id).first()
        if not content:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="内容不存在"
            )

        # 使用共享的内容分析服务进行智能分析
        analysis_service = get_content_analysis_service_instance()

        # 使用增强的智能分析
        result = await analysis_service.analyze_content_smart(
            content=content,
            content_type=content.content_type or "text",
            analysis_type=analysis_type,
            use_multimodal=True
        )

        if not result.get("success"):
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"内容分析失败: {result.get('error', '未知错误')}"
            )

        # 更新数据库记录
        updated_fields = {}

        if analysis_type in ['summary', 'both'] and result.get('ai_summary'):
            # 只有在摘要未被用户手动修改过时才更新
            if not hasattr(content, 'summary_manually_modified') or not content.summary_manually_modified:
                content.content_summary = result['ai_summary']
                updated_fields['summary'] = result['ai_summary']

        if analysis_type in ['tags', 'both'] and result.get('ai_tags'):
            # 只有在标签未被用户手动修改过时才更新
            if not hasattr(content, 'tags_manually_modified') or not content.tags_manually_modified:
                tags = result['ai_tags']
                if isinstance(tags, list):
                    # 过滤无效标签
                    valid_tags = [tag.strip() for tag in tags if isinstance(tag, str) and tag.strip()]
                    content.content_tags = valid_tags
                    updated_fields['tags'] = valid_tags

        # 更新元数据信息
        if hasattr(content, 'content_metadata') and content.content_metadata:
            try:
                metadata = json.loads(content.content_metadata) if isinstance(content.content_metadata, str) else content.content_metadata
            except:
                metadata = {}
        else:
            metadata = {}

        # 添加分析信息到元数据
        metadata.update({
            "last_analysis_time": datetime.now().isoformat(),
            "analysis_method": result.get("analysis_method", "enhanced_service"),
            "analysis_confidence": result.get("confidence", 0.8),
            "content_category": result.get("content_category", ""),
            "key_points_count": len(result.get("key_points", []))
        })

        content.content_metadata = json.dumps(metadata, ensure_ascii=False)

        db.commit()
        db.refresh(content)

        # 更新 doc-profile 向量（仅在摘要/标签发生更新时）
        if updated_fields:
            try:
                rag_service = get_enhanced_rag_service()
                await rag_service.upsert_doc_profile_vector(content)
            except Exception as e:
                logger.warning(f"更新文档概要向量失败（已忽略）: {e}")

        return {
            "success": True,
            "message": "增强内容分析完成",
            "data": {
                "content_id": content_id,
                "summary": result.get('ai_summary'),
                "tags": result.get('ai_tags'),
                "key_points": result.get('key_points', []),
                "content_category": result.get('content_category'),
                "confidence": result.get('confidence'),
                "analysis_method": result.get('analysis_method'),
                "analysis_type": analysis_type,
                "updated_fields": list(updated_fields.keys())
            }
        }

    except HTTPException:
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"单个内容分析失败: {str(e)}"
        )


async def _get_content_by_resource_id(db: Session, resource_id: str) -> Optional[RAGContent]:
    """按 resource_id 获取最新内容记录。"""
    return db.query(RAGContent) \
        .filter(RAGContent.resource_id == resource_id) \
        .order_by(RAGContent.update_time.desc()) \
        .first()


async def _analyze_single_resource(resource_id: str, analysis_type: str, db: Session) -> Dict:
    """按资源ID分析单个内容。"""
    try:
        content = await _get_content_by_resource_id(db, resource_id)
        if not content:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="资源不存在或尚未生成内容")

        analysis_service = get_content_analysis_service_instance()
        result = await analysis_service.analyze_content_smart(
            content=content,
            content_type=content.content_type or "text",
            analysis_type=analysis_type,
            use_multimodal=True
        )

        if not result.get("success"):
            raise HTTPException(status_code=500, detail=f"内容分析失败: {result.get('error', '未知错误')}")

        # 根据返回更新数据库（沿用原逻辑）
        updated_fields = {}
        if analysis_type in ['summary', 'both'] and result.get('ai_summary'):
            if not hasattr(content, 'summary_manually_modified') or not content.summary_manually_modified:
                content.content_summary = result['ai_summary']
                updated_fields['summary'] = result['ai_summary']
        if analysis_type in ['tags', 'both'] and result.get('ai_tags'):
            if not hasattr(content, 'tags_manually_modified') or not content.tags_manually_modified:
                tags = result['ai_tags']
                if isinstance(tags, list):
                    valid_tags = [t.strip() for t in tags if isinstance(t, str) and t.strip()]
                    content.content_tags = valid_tags
                    updated_fields['tags'] = valid_tags

        # 元数据更新
        try:
            metadata = json.loads(content.content_metadata) if isinstance(content.content_metadata, str) else (content.content_metadata or {})
        except Exception:
            metadata = {}
        metadata.update({
            "last_analysis_time": datetime.now().isoformat(),
            "analysis_method": result.get("analysis_method", "enhanced_service"),
            "analysis_confidence": result.get("confidence", 0.8),
            "content_category": result.get("content_category", ""),
            "key_points_count": len(result.get("key_points", []))
        })
        content.content_metadata = json.dumps(metadata, ensure_ascii=False)
        db.commit()
        db.refresh(content)

        return {
            "success": True,
            "message": "增强内容分析完成",
            "data": {
                "resource_id": resource_id,
                "content_id": content.id,
                "summary": result.get('ai_summary'),
                "tags": result.get('ai_tags'),
                "key_points": result.get('key_points', []),
                "content_category": result.get('content_category'),
                "confidence": result.get('confidence'),
                "analysis_method": result.get('analysis_method'),
                "analysis_type": analysis_type,
                "updated_fields": list(updated_fields.keys())
            }
        }
    except HTTPException:
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"单个资源分析失败: {str(e)}")


async def _analyze_batch_resource(resource_ids: List[str], analysis_type: str, db: Session) -> Dict:
    """按资源ID批量分析。"""
    try:
        analysis_service = get_content_analysis_service_instance()
        content_list = []
        missing_ids: List[str] = []

        for rid in resource_ids:
            content = await _get_content_by_resource_id(db, rid)
            if content:
                text_content = content.content_text or ""
                if not text_content and content.file_path:
                    try:
                        text_content = await _extract_text_from_file(content.file_path, content.file_name)
                    except Exception:
                        text_content = f"文档: {content.file_name}"
                content_list.append({
                    "id": content.id,  # 保持与分析器接口的兼容（内部仍用content.id）
                    "resource_id": rid,
                    "content": text_content,
                    "content_type": content.content_type or "text",
                    "analysis_type": analysis_type
                })
            else:
                missing_ids.append(rid)

        logger.info(f"[ContentAnalysis] 资源可分析数: {len(content_list)}, 缺失: {missing_ids}")

        # 使用现有批量分析逻辑（传 content_list）
        batch_result = await analysis_service.batch_analyze_contents(content_list, analysis_type)

        results = []
        processed_content_ids = set()

        if batch_result.get("success"):
            for item in batch_result.get("results", []):
                cid = item.get("item_id")
                if not cid:
                    continue
                processed_content_ids.add(cid)
                # 找到对应的资源ID
                rid = next((x["resource_id"] for x in content_list if x["id"] == cid), None)
                if item.get("success"):
                    data = item.get("data", {})
                    # 更新数据库字段（沿用旧逻辑），按 content.id 定位
                    content = db.query(RAGContent).filter(RAGContent.id == cid).first()
                    if content:
                        updated_fields = []
                        if analysis_type in ['summary', 'both'] and data.get('summary'):
                            if not hasattr(content, 'summary_manually_modified') or not content.summary_manually_modified:
                                content.content_summary = data['summary']
                                updated_fields.append('summary')
                        if analysis_type in ['tags', 'both'] and data.get('tags'):
                            if not hasattr(content, 'tags_manually_modified') or not content.tags_manually_modified:
                                tags = data['tags']
                                if isinstance(tags, list):
                                    valid_tags = [t.strip() for t in tags if isinstance(t, str) and t.strip()]
                                    content.content_tags = valid_tags
                                    updated_fields.append('tags')
                        db.commit()

                        if updated_fields:
                            try:
                                rag_service = get_enhanced_rag_service()
                                await rag_service.upsert_doc_profile_vector(content)
                            except Exception as e:
                                logger.warning(f"更新文档概要向量失败（已忽略）: {e}")

                        results.append({
                            "resource_id": rid,
                            "content_id": cid,
                            "success": True,
                            "summary": data.get('summary'),
                            "tags": data.get('tags', []),
                            "key_points": data.get('key_points', []),
                            "content_category": data.get('content_category'),
                            "confidence": data.get('confidence'),
                            "updated_fields": updated_fields
                        })
                    else:
                        results.append({
                            "resource_id": rid,
                            "content_id": cid,
                            "success": False,
                            "error": "内容不存在"
                        })
                else:
                    results.append({
                        "resource_id": rid,
                        "success": False,
                        "error": item.get("error", "分析失败")
                    })

        # 未处理的（在content_list中但未返回）
        for item in content_list:
            cid = item["id"]
            if cid not in processed_content_ids:
                results.append({
                    "resource_id": item.get("resource_id"),
                    "content_id": cid,
                    "success": False,
                    "error": "Agent未返回此内容的分析结果"
                })

        # 缺失的资源ID
        for rid in missing_ids:
            results.append({
                "resource_id": rid,
                "success": False,
                "error": "资源不存在或尚未生成内容"
            })

        success_count = len([r for r in results if r.get('success')])
        failed_count = len(results) - success_count

        return {
            "success": True,
            "message": f"批量分析完成，成功: {success_count}，失败: {failed_count}",
            "data": results,
            "meta": {
                "total": len(resource_ids),
                "success_count": success_count,
                "failed_count": failed_count
            }
        }
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"批量资源分析失败: {str(e)}")


async def _analyze_batch_content(content_ids: List[str], analysis_type: str, db: Session) -> Dict:
    """分析批量内容 - 使用增强的共享分析服务"""
    try:
        # 获取content analysis service
        analysis_service = get_content_analysis_service_instance()

        # 准备批量分析的内容列表
        content_list = []
        missing_content_ids = []

        for content_id in content_ids:
            content = db.query(RAGContent).filter(RAGContent.id == content_id).first()
            if content:
                # 获取文本内容
                text_content = content.content_text or ""
                if not text_content and content.file_path:
                    try:
                        text_content = await _extract_text_from_file(content.file_path, content.file_name)
                    except:
                        text_content = f"文档: {content.file_name}"

                content_list.append({
                    "id": content_id,
                    "content": text_content,
                    "content_type": content.content_type or "text",
                    "analysis_type": analysis_type
                })
            else:
                # 记录找不到的content_id
                missing_content_ids.append(content_id)

        logger.info(f"准备分析的内容数量: {len(content_list)}")
        logger.info(f"找不到的内容ID: {missing_content_ids}")

        # 初始化结果列表
        results = []

        # 先为找不到的内容添加失败记录
        for content_id in missing_content_ids:
            results.append({
                "content_id": content_id,
                "success": False,
                "error": "内容不存在"
            })

        # 如果有可分析的内容，调用增强服务进行批量分析
        if content_list:
            batch_result = await analysis_service.batch_analyze_contents(content_list, analysis_type)
            logger.info(f"增强服务返回结果数量: {len(batch_result.get('results', []))}")

            # 处理增强服务返回的结果
            processed_ids = set()
            for item in batch_result.get("results", []):
                content_id = item.get("item_id")
                if content_id:
                    processed_ids.add(content_id)

                    if item.get("success"):
                        try:
                            content = db.query(RAGContent).filter(RAGContent.id == content_id).first()
                            if content:
                                data = item.get("data", {})
                                updated_fields = []

                                # 更新摘要
                                if 'summary' in data and data['summary']:
                                    # 只有在摘要未被用户手动修改过时才更新
                                    if not hasattr(content, 'summary_manually_modified') or not content.summary_manually_modified:
                                        content.content_summary = data['summary']
                                        updated_fields.append('summary')

                                # 更新标签，确保是有效的JSON数组
                                if 'tags' in data and data['tags']:
                                    # 只有在标签未被用户手动修改过时才更新
                                    if not hasattr(content, 'tags_manually_modified') or not content.tags_manually_modified:
                                        # 确保tags是列表格式
                                        tags = data['tags']
                                        if isinstance(tags, list):
                                            # 过滤掉无效的标签
                                            valid_tags = [tag.strip() for tag in tags if isinstance(tag, str) and tag.strip()]
                                            content.content_tags = valid_tags
                                            updated_fields.append('tags')
                                        else:
                                            content.content_tags = [str(tags)]
                                            updated_fields.append('tags')

                                # 更新元数据信息
                                if hasattr(content, 'content_metadata') and content.content_metadata:
                                    try:
                                        metadata = json.loads(content.content_metadata) if isinstance(content.content_metadata, str) else content.content_metadata
                                    except:
                                        metadata = {}
                                else:
                                    metadata = {}

                                # 添加分析信息到元数据
                                metadata.update({
                                    "last_analysis_time": datetime.now().isoformat(),
                                    "batch_analysis": True,
                                    "analysis_confidence": data.get("confidence", 0.8),
                                    "content_category": data.get("content_category", ""),
                                    "key_points_count": len(data.get("key_points", []))
                                })

                                content.content_metadata = json.dumps(metadata, ensure_ascii=False)

                                db.commit()

                                if updated_fields:
                                    try:
                                        rag_service = get_enhanced_rag_service()
                                        await rag_service.upsert_doc_profile_vector(content)
                                    except Exception as e:
                                        logger.warning(f"更新文档概要向量失败（已忽略）: {e}")

                                # 获取清理后的标签用于返回
                                return_tags = content.content_tags if content.content_tags else []

                                results.append({
                                    "content_id": content_id,
                                    "success": True,
                                    "summary": data.get('summary'),
                                    "tags": return_tags,
                                    "key_points": data.get('key_points', []),
                                    "content_category": data.get('content_category'),
                                    "confidence": data.get('confidence'),
                                    "updated_fields": updated_fields
                                })
                            else:
                                results.append({
                                    "content_id": content_id,
                                    "success": False,
                                    "error": "内容不存在"
                                })

                        except Exception as e:
                            db.rollback()
                            results.append({
                                "content_id": content_id,
                                "success": False,
                                "error": str(e)
                            })
                    else:
                        results.append({
                            "content_id": content_id,
                            "success": False,
                            "error": item.get("error", "分析失败")
                        })

            # 检查是否有发送给agent但未被处理的content_id
            content_list_ids = {item["id"] for item in content_list}
            for content_id in content_list_ids:
                if content_id not in processed_ids:
                    results.append({
                        "content_id": content_id,
                        "success": False,
                        "error": "Agent未返回此内容的分析结果"
                    })

        success_count = len([r for r in results if r['success']])
        failed_count = len(results) - success_count

        return {
            "success": True,
            "message": f"批量分析完成，成功: {success_count}，失败: {failed_count}",
            "data": results,  # 直接返回结果数组
            "meta": {
                "total": len(content_ids),
                "success_count": success_count,
                "failed_count": failed_count
            }
        }

    except Exception as e:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"批量分析失败: {str(e)}"
        )



@router.post("/batch-analyze", response_model=Dict)
async def batch_analyze_content(
    content_ids: List[str],
    analysis_type: str = "both",
    db: Session = Depends(get_db)
):
    """
    批量分析内容 - 使用增强的分析服务
    """
    try:
        # 直接使用新的批量分析函数
        return await _analyze_batch_content(content_ids, analysis_type, db)

    except Exception as e:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"批量分析失败: {str(e)}"
        )


@router.post("/get-result", response_model=Dict)
async def get_content_analysis(
    request: dict,
    db: Session = Depends(get_db)
):
    """
    获取内容的摘要和标签
    """
    try:
        # 获取内容ID
        content_id = request.get("content_id")
        if not content_id:
            raise HTTPException(status_code=400, detail="缺少必要参数: content_id")

        content = db.query(RAGContent).filter(RAGContent.id == content_id).first()
        if not content:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="内容不存在"
            )

        tags = []
        if content.content_tags:
            try:
                tags = json.loads(content.content_tags)
            except:
                tags = []

        return {
            "success": True,
            "data": {
                "content_id": content_id,
                "title": content.title,
                "content_type": content.content_type,
                "summary": content.content_summary,
                "tags": tags,
                "file_name": content.file_name
            }
        }

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"获取分析结果失败: {str(e)}"
        )
