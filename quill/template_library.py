"""Original, non-customer starter templates, rendered to real editable DOCX files."""

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def starter(key, name, category, description, fields, sections):
    return dict(
        id=key,
        name=name,
        category=category,
        description=description,
        fields=fields,
        sections=[dict(title=title, field=field) for title, field in sections],
    )


TEMPLATES = [
    starter(
        "project-plan",
        "项目实施方案",
        "方案",
        "从项目目标到交付验收，组织完整的实施计划。",
        ["项目名称", "编制单位", "负责人", "编制日期"],
        [
            ("项目背景与目标", "项目背景与目标"),
            ("实施范围", "实施范围"),
            ("实施步骤与里程碑", "实施步骤与里程碑"),
            ("人员与资源", "人员与资源"),
            ("风险与应对", "风险与应对"),
            ("交付与验收", "交付与验收"),
        ],
    ),
    starter(
        "technical-plan",
        "技术方案",
        "方案",
        "围绕需求、架构、接口与部署说明技术设计。",
        ["方案名称", "编制单位", "技术负责人", "版本日期"],
        [
            ("需求与约束", "需求与约束"),
            ("总体架构", "总体架构"),
            ("功能设计", "功能设计"),
            ("数据与接口", "数据与接口"),
            ("部署与安全", "部署与安全"),
            ("测试与验证", "测试与验证"),
        ],
    ),
    starter(
        "operations-plan",
        "运维服务方案",
        "方案",
        "明确巡检、监控、故障处理与服务交付安排。",
        ["项目名称", "服务单位", "服务负责人", "服务周期"],
        [
            ("服务范围", "服务范围"),
            ("巡检与监控", "巡检与监控"),
            ("故障响应流程", "故障响应流程"),
            ("备份与恢复", "备份与恢复"),
            ("服务报告与交付", "服务报告与交付"),
        ],
    ),
    starter(
        "acceptance-report",
        "项目验收报告",
        "报告",
        "核对交付清单、测试结果与验收意见，留存签字记录。",
        ["项目名称", "建设单位", "交付单位", "验收日期"],
        [
            ("项目概况", "项目概况"),
            ("交付清单", "交付清单"),
            ("测试与核查结果", "测试与核查结果"),
            ("遗留问题与整改", "遗留问题与整改"),
            ("验收意见", "验收意见"),
            ("签字确认", "签字确认"),
        ],
    ),
    starter(
        "weekly-report",
        "工作周报",
        "报告",
        "记录本周成果、问题和下周计划，适合个人与团队复用。",
        ["团队或项目", "汇报人", "报告周期", "提交日期"],
        [
            ("本周完成", "本周完成"),
            ("进度与成果", "进度与成果"),
            ("问题与需协助事项", "问题与需协助事项"),
            ("下周计划", "下周计划"),
        ],
    ),
    starter(
        "meeting-minutes",
        "会议纪要",
        "报告",
        "整理议题、讨论结论与行动项，便于会后跟进。",
        ["会议主题", "会议时间", "会议地点", "参会人员", "主持人", "记录人"],
        [
            ("会议议题", "会议议题"),
            ("讨论与结论", "讨论与结论"),
            ("行动项与负责人", "行动项与负责人"),
            ("完成时限与跟进", "完成时限与跟进"),
        ],
    ),
    starter(
        "progress-report",
        "项目进展报告",
        "报告",
        "同步里程碑、资源投入、风险和下一阶段安排。",
        ["项目名称", "报告人", "报告周期", "报告日期"],
        [
            ("总体进展", "总体进展"),
            ("里程碑完成情况", "里程碑完成情况"),
            ("资源投入", "资源投入"),
            ("风险与问题", "风险与问题"),
            ("下一阶段计划", "下一阶段计划"),
        ],
    ),
    starter(
        "quotation",
        "商务报价单",
        "其他",
        "汇总报价明细、交付范围、有效期与商务说明。",
        ["报价名称", "报价单位", "联系人", "报价日期", "客户名称", "报价有效期"],
        [
            ("需求与报价范围", "需求与报价范围"),
            ("报价明细", "报价明细"),
            ("合计与税费说明", "合计与税费说明"),
            ("交付与付款安排", "交付与付款安排"),
            ("备注", "备注"),
        ],
    ),
    starter(
        "bid-response",
        "投标响应文件",
        "招投标",
        "提供投标文件组织框架，具体条款按招标要求核对补齐。",
        ["项目名称", "项目编号", "投标单位", "授权代表", "编制日期"],
        [
            ("投标说明", "投标说明"),
            ("单位与资质介绍", "单位与资质介绍"),
            ("商务响应", "商务响应"),
            ("技术响应", "技术响应"),
            ("实施与服务承诺", "实施与服务承诺"),
            ("附件清单", "附件清单"),
        ],
    ),
    starter(
        "resume",
        "个人简历",
        "简历",
        "按个人简介、工作经历与项目经验组织求职资料。",
        ["姓名", "意向岗位", "联系电话", "电子邮箱"],
        [
            ("个人简介", "个人简介"),
            ("教育背景", "教育背景"),
            ("工作经历", "工作经历"),
            ("项目经验", "项目经验"),
            ("技能与证书", "技能与证书"),
        ],
    ),
    starter(
        "business-letter",
        "商务函件",
        "函件",
        "适用于沟通、情况说明与事项确认，保留正式函件结构。",
        ["函件主题", "收件单位", "发件单位", "联系人", "发函日期"],
        [
            ("事由与背景", "事由与背景"),
            ("正文", "正文"),
            ("需确认或办理事项", "需确认或办理事项"),
            ("附件与落款", "附件与落款"),
        ],
    ),
    starter(
        "requirements",
        "需求说明书",
        "方案",
        "梳理业务流程、功能需求、边界与验收标准。",
        ["产品或项目", "需求负责人", "版本号", "更新日期"],
        [
            ("背景与目标", "背景与目标"),
            ("使用者与业务流程", "使用者与业务流程"),
            ("功能需求", "功能需求"),
            ("非功能需求", "非功能需求"),
            ("范围与边界", "范围与边界"),
            ("验收标准", "验收标准"),
        ],
    ),
]


def catalog() -> list[dict]:
    return [
        dict(
            **template,
            origin="builtin",
            variables=[
                dict(name=name, default_value="")
                for name in template["fields"] + [s["field"] for s in template["sections"]]
            ],
        )
        for template in TEMPLATES
    ]


def get_template(key: str) -> dict | None:
    return next((template for template in catalog() if template["id"] == key), None)


def render_starter(template: dict) -> bytes:
    document = Document()
    document.core_properties.title = template["name"]
    document.core_properties.author = "Otterview"
    document.core_properties.language = "zh-CN"
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.5)
    for name, size in [("Normal", 11), ("Title", 24), ("Heading 1", 14)]:
        style = document.styles[name]
        style.font.name = "Calibri"
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string("203D2B" if name == "Heading 1" else "202923")
    document.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    for border in document.styles["Title"].element.xpath("./w:pPr/w:pBdr"):
        border.getparent().remove(border)
    normal = document.styles["Normal"].paragraph_format
    normal.line_spacing = 1.4
    normal.space_after = Pt(8)
    heading = document.styles["Heading 1"].paragraph_format
    heading.space_before, heading.space_after = Pt(16), Pt(8)
    section.header.paragraphs[0].text = template["name"]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页")
    document.add_heading(template["name"], 0)
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    table.columns[0].width, table.columns[1].width = Cm(4), Cm(12)
    for name in template["fields"]:
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = "{{ " + name + " }}"
        cells[0].width, cells[1].width = Cm(4), Cm(12)
    for item in template["sections"]:
        document.add_heading(item["title"], 1)
        document.add_paragraph("{{ " + item["field"] + " }}")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()
