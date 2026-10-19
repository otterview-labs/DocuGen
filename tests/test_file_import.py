import io

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from quill.config import Settings
from quill.file_import import ImportError, extract_text
from quill.main import create_app
from quill.providers import KnowledgeClient, ProviderError


def test_word_import_keeps_table_order():
    doc = Document()
    doc.add_paragraph("项目背景")
    doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "验收范围"
    doc.tables[0].cell(0, 1).text = "设备接入"
    doc.add_paragraph("交付清单")
    data = io.BytesIO()
    doc.save(data)
    text = extract_text(data.getvalue(), "方案.docx")
    assert text.index("项目背景") < text.index("验收范围") < text.index("交付清单")
    assert "设备接入" in text


def test_pdf_text_extraction_and_scanned_pdf_error():
    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    data = io.BytesIO()
    writer.write(data)
    with pytest.raises(ImportError, match="OCR"):
        extract_text(data.getvalue(), "scanned.pdf")
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 10 180 Td (Acceptance period: 30 days) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    data = io.BytesIO()
    writer.write(data)
    assert "30 days" in extract_text(data.getvalue(), "source.pdf")


@pytest.mark.parametrize("filename", ["source.md", "source.txt", "source.csv", "source.json"])
def test_text_formats_and_chinese_encoding(filename):
    assert extract_text("知识库资料\r\n验收说明".encode("gb18030"), filename) == "知识库资料\n验收说明"


@pytest.mark.parametrize(
    "data,filename",
    [(b"", "empty.txt"), (b"bad", "bad.docx"), (b"fake", "unsupported.exe"), (b"a\x00b", "binary.txt")],
)
def test_rejects_invalid_files(data, filename):
    with pytest.raises(ImportError):
        extract_text(data, filename)


def test_file_upload_duplicate_original_download_and_delete(tmp_path):
    class Knowledge:
        def __init__(self):
            self.bodies = []

        async def request(self, path, body):
            if path == "quill/text":
                self.bodies.append(body)
                return {
                    "result": {
                        "resource_id": "r1",
                        "vector_indexed": False,
                        "duplicate": len(self.bodies) > 1,
                    }
                }
            if path == "quill/resources/get":
                return {"result": {"resource_id": "r1", "content": "验收资料", "segments": []}}
            if path == "quill/resources/delete":
                return {"result": {"deleted": True}}
            raise AssertionError(path)

    config = Settings(data_dir=tmp_path, api_key="test", _env_file=None)
    with TestClient(create_app(config)) as client:
        client.headers["Authorization"] = "Bearer test"
        provider = Knowledge()
        client.app.state.knowledge = provider
        for _ in range(2):
            response = client.post(
                "/api/v1/knowledge/import",
                data={"collection_id": "c1"},
                files={"file": ("../原文件.txt", "验收资料".encode())},
            )
            assert response.status_code == 200, response.text
        assert response.json()["duplicate"] is True
        assert provider.bodies[0]["file_name"] == "原文件.txt"
        assert len(provider.bodies[0]["file_sha256"]) == 64
        assert len(list((tmp_path / "knowledge-files").iterdir())) == 1
        assert client.get("/api/v1/knowledge/resources/r1").json()["has_original"] is True
        response = client.get("/api/v1/knowledge/resources/r1/download")
        assert response.content.decode() == "验收资料"
        assert (
            client.get("/api/v1/knowledge/resources/r1/download", headers={"Authorization": ""}).status_code
            == 401
        )
        assert client.delete("/api/v1/knowledge/resources/r1").status_code == 200
        assert not list((tmp_path / "knowledge-files").iterdir())


@pytest.mark.parametrize(
    "body",
    [{"detail": "向量模型尚未就绪"}, {"error": True, "message": "向量模型尚未就绪", "status_code": 409}],
)
@pytest.mark.parametrize("status", [409, 502])
async def test_unconfigured_index_error_retains_http_status(body, status):
    async with httpx.AsyncClient(
        transport=httpx.MockTransport(lambda req: httpx.Response(status, json=body))
    ) as client:
        provider = KnowledgeClient(Settings(_env_file=None), client)
        with pytest.raises(ProviderError) as failure:
            await provider.request("quill/resources/reindex", {"resource_id": "r1"})
        assert failure.value.status_code == status
        assert str(failure.value) == "向量模型尚未就绪"
