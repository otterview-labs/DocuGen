import io
import json

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient

from quill.config import Settings
from quill.main import create_app
from quill.providers import KnowledgeClient, LLMClient, ProviderError
from quill.store import Store


@pytest.fixture
def client(tmp_path):
    settings = Settings(data_dir=tmp_path, api_key="test-key", llm_model="test-model", _env_file=None)
    with TestClient(create_app(settings)) as client:
        client.headers["Authorization"] = "Bearer test-key"
        yield client


def upload(client, name="方案"):
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("项目：{{项").bold = True
    p.add_run("目名称}}；负责人：【负责人】")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "{{项目名称}}"
    data = io.BytesIO()
    doc.save(data)
    return client.post(
        "/api/upload-template",
        data={"template_name": name},
        files={"file": ("template.docx", data.getvalue())},
    )


def test_auth_and_unknown_api(client):
    assert client.get("/health", headers={"Authorization": ""}).status_code == 200
    assert client.get("/api/templates", headers={"Authorization": ""}).status_code == 401
    assert client.get("/api/missing").status_code == 404


def test_startup_with_environment_socks_proxy(tmp_path, monkeypatch):
    # Initializing HTTPX must work with a configured proxy even before any
    # external request. This previously crashed without the SOCKS dependency.
    monkeypatch.setenv("ALL_PROXY", "socks5://127.0.0.1:9")
    settings = Settings(data_dir=tmp_path, api_key="", _env_file=None)
    with TestClient(create_app(settings)) as local:
        assert local.get("/health").status_code == 200
        assert local.get("/api/templates").json()["templates"] == []


@pytest.mark.parametrize("access_key", ["", "test-key"])
def test_optional_workspace_access_covers_upload_render_and_download(tmp_path, access_key):
    settings = Settings(data_dir=tmp_path, api_key=access_key, _env_file=None)
    with TestClient(create_app(settings)) as local:
        if access_key:
            assert local.get("/api/v1/status").status_code == 401
            assert local.get("/api/templates").status_code == 401
            local.headers["Authorization"] = f"Bearer {access_key}"
        status = local.get("/api/v1/status")
        assert status.status_code == 200
        assert status.json()["authentication"] is bool(access_key)
        assert upload(local, "示例 · 访问模式验收").status_code == 200
        rendered = local.post(
            "/api/generate-from-template",
            data={
                "template_name": "示例 · 访问模式验收",
                "data": json.dumps({"项目名称": "本地验收", "负责人": "示例人员"}),
            },
        )
        assert rendered.status_code == 200
        url = "/download/" + rendered.json()["filename"]
        download = local.get(url)
        assert download.status_code == 200
        document = Document(io.BytesIO(download.content))
        assert "本地验收" in document.paragraphs[0].text
        assert "示例人员" in document.paragraphs[0].text
        if access_key:
            assert local.get(url, headers={"Authorization": ""}).status_code == 401
            assert local.get(url, headers={"Authorization": "Bearer wrong-key"}).status_code == 401
        else:
            # An old tab's saved token must not prevent key-free local use.
            assert local.get(url, headers={"Authorization": "Bearer old-key"}).status_code == 200


def test_template_render_roundtrip_and_persistence(client):
    response = upload(client)
    assert response.status_code == 200, response.text
    assert {v["name"] for v in response.json()["variables"]} == {"项目名称", "负责人"}
    assert upload(client).status_code == 409
    template = client.get("/api/templates").json()["templates"][0]
    assert template["name"] == "方案"
    sources = [{"title": "项目事实", "resource_id": "res-1"}]
    result = client.post(
        "/api/generate-from-template",
        data={
            "template_name": "方案",
            "data": json.dumps({"项目名称": "测试工程", "负责人": "张三"}),
            "sources": json.dumps(sources),
        },
    )
    assert result.status_code == 200, result.text
    filename = result.json()["filename"]
    download = client.get(f"/download/{filename}")
    doc = Document(io.BytesIO(download.content))
    assert "测试工程" in doc.paragraphs[0].text
    assert "张三" in doc.paragraphs[0].text
    assert doc.tables[0].cell(0, 0).text == "测试工程"
    assert doc.paragraphs[0].runs[0].bold is True
    assert client.get(f"/api/v1/artifacts/{filename}/sources").json()["sources"] == sources
    reopened = Store(client.app.state.store.root)
    assert reopened.template("方案")["id"] == template["id"]
    assert reopened.artifact(filename)["sources"] == sources
    assert client.get("/download/not-a-registered-file.docx").status_code == 404


def test_template_headers_footers_nested_tables_and_fields_roundtrip(client):
    from docx.oxml import OxmlElement

    document = Document()
    document.add_paragraph("项目：{{项目名称}}")
    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.add_run("报告：{{页").bold = True
    header.add_run("眉名称}}")
    section.first_page_header.paragraphs[0].text = "封面：【封面名称】"
    section.even_page_header.paragraphs[0].text = "偶数页：{{偶数页名称}}"
    footer = section.footer.paragraphs[0]
    footer.add_run("编制：【编制人】 页码：")
    field = OxmlElement("w:fldSimple")
    field.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}instr", "PAGE")
    footer._p.append(field)
    table = document.add_table(rows=1, cols=2)
    cell = table.cell(0, 0).merge(table.cell(0, 1))
    cell.text = "外表：{{外表值}}"
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "内表：【内表值】"
    data = io.BytesIO()
    document.save(data)
    uploaded = client.post(
        "/api/upload-template",
        data={"template_name": "示例 · 复杂位置变量"},
        files={"file": ("stories.docx", data.getvalue())},
    )
    assert uploaded.status_code == 200
    values = {"项目名称": "本地项目", "页眉名称": "项目报告", "编制人": "示例甲",
              "外表值": "外层", "内表值": "内层", "封面名称": "验收封面", "偶数页名称": "附录"}
    assert {item["name"] for item in uploaded.json()["variables"]} == set(values)
    assert len(uploaded.json()["variables"]) == len(values)
    renderer = client.app.state.store
    metadata = renderer.template("示例 · 复杂位置变量")
    from quill.legacy.template import BaseTemplatesGenerator
    text = BaseTemplatesGenerator(renderer.root / "downloads").get_template_content(
        str(renderer.root / "templates" / (metadata["id"] + ".docx"))
    )
    assert "页眉名称" in text and "内表值" in text and "编制人" in text
    rendered = client.post("/api/generate-from-template", data={
        "template_name": "示例 · 复杂位置变量", "data": json.dumps(values),
    })
    assert rendered.status_code == 200
    result = Document(io.BytesIO(client.get("/download/" + rendered.json()["filename"]).content))
    assert result.sections[0].header.paragraphs[0].text == "报告：项目报告"
    assert result.sections[0].header.paragraphs[0].runs[0].bold is True
    assert result.sections[0].first_page_header.paragraphs[0].text == "封面：验收封面"
    assert result.sections[0].even_page_header.paragraphs[0].text == "偶数页：附录"
    assert "编制：示例甲" in result.sections[0].footer.paragraphs[0].text
    assert result.sections[0].footer._element.xpath(".//w:fldSimple")[0].get(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}instr") == "PAGE"
    assert "外层" in result.tables[0].cell(0, 0).text
    assert result.tables[0].cell(0, 0).tables[0].cell(0, 0).text == "内表：内层"


def test_invalid_upload_and_values(client):
    assert client.post("/api/upload-template", files={"file": ("x.doc", b"bad")}).status_code == 400
    assert client.post("/api/upload-template", files={"file": ("x.docx", b"bad")}).status_code == 400
    upload(client)
    assert (
        client.post("/api/generate-from-template", data={"template_name": "方案", "data": "[]"}).status_code
        == 400
    )
    assert (
        client.post(
            "/api/ai-fill-template-variables", json={"template_name": "方案", "enable_rag": True}
        ).status_code
        == 400
    )


async def test_knowledge_contract_and_pagination():
    seen = []

    def handle(request):
        seen.append(json.loads(request.content))
        assert request.headers["Authorization"] == "Bearer engine-token"
        if request.url.path.endswith("collections/search"):
            page = seen[-1]["pageIndex"]
            if page == 0:
                return httpx.Response(
                    200,
                    json={
                        "result": {"data": [{"id": str(i), "name": str(i)} for i in range(100)], "total": 101}
                    },
                )
            return httpx.Response(
                200, json={"result": {"data": [{"id": "last", "name": "最后一个"}], "total": 101}}
            )
        assert seen[-1]["collection_id"] == ["last"]
        return httpx.Response(
            200,
            json={
                "results": [
                    {
                        "resource_id": "r1",
                        "content_id": "c1",
                        "content": "参考事实",
                        "title": "技术资料",
                        "collection_id": "last",
                        "version": 2,
                    }
                ]
            },
        )

    async with httpx.AsyncClient(transport=httpx.MockTransport(handle)) as client:
        provider = KnowledgeClient(Settings(knowledge_token="engine-token", _env_file=None), client)
        assert len(await provider.collections()) == 101
        sources = await provider.search("事实", ["last"])
        assert sources[0]["resource_id"] == "r1"
        assert sources[0]["version"] == 2


async def test_provider_failure_never_becomes_empty_success():
    async with httpx.AsyncClient(transport=httpx.MockTransport(lambda req: httpx.Response(503))) as client:
        provider = KnowledgeClient(Settings(_env_file=None), client)
        with pytest.raises(ProviderError):
            await provider.collections()
        llm = LLMClient(Settings(llm_model="model", _env_file=None), client)
        with pytest.raises(ProviderError):
            await llm.chat("system", "prompt")


@pytest.mark.parametrize("effort", ["", "none"])
async def test_keyless_model_contract_and_optional_reasoning(effort):
    def respond(request):
        assert "Authorization" not in request.headers
        body = json.loads(request.content)
        if effort:
            assert body["reasoning_effort"] == effort
        else:
            assert "reasoning_effort" not in body
        return httpx.Response(200, json={"choices": [{"message": {"content": "本机模型验收成功"}}]})

    settings = Settings(llm_model="local-model", llm_reasoning_effort=effort, _env_file=None)
    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        assert await LLMClient(settings, client).chat("验收", "测试") == "本机模型验收成功"


async def test_model_reasoning_without_a_final_answer_is_not_success():
    response = {"choices": [{"message": {"content": "", "reasoning": "尚未完成的推理"}}]}
    async with httpx.AsyncClient(transport=httpx.MockTransport(lambda r: httpx.Response(200, json=response))) as c:
        with pytest.raises(ProviderError):
            await LLMClient(Settings(llm_model="local-model", _env_file=None), c).chat("验收", "测试")


async def test_generation_sections_sources_and_failure(tmp_path):
    from quill.generation import GenerationService

    class Knowledge:
        async def search(self, query, collections):
            return [{"title": "证据", "content": "工期30天", "resource_id": "r1"}]

    class Model:
        async def chat(self, system, prompt):
            return "根据资料，工期为30天 [1]。"

    store = Store(tmp_path)
    service = GenerationService(store, Knowledge(), Model())
    job = store.new_job({"topic": "实施方案", "sections": ["背景", "计划"], "collection_ids": ["c1"]})
    await service.run(job["id"])
    result = store.job(job["id"])
    assert result["status"] == "completed"
    assert result["progress"] == 100
    assert [s["section"] for s in result["sources"]] == ["背景", "计划"]
    assert (tmp_path / "downloads" / result["filename"]).is_file()
    assert store.artifact(result["filename"])["job_id"] == job["id"]

    class EmptyKnowledge:
        async def search(self, query, collections):
            return []

    service.knowledge = EmptyKnowledge()
    failed = store.new_job({"topic": "实施方案", "collection_ids": ["c1"]})
    await service.run(failed["id"])
    assert store.job(failed["id"])["status"] == "failed"
    assert "没有检索到" in store.job(failed["id"])["error"]


def test_restart_recovers_all_pending_jobs(tmp_path):
    store = Store(tmp_path)
    one = store.new_job({"topic": "one"})
    two = store.new_job({"topic": "two"})
    two.update(status="running", content="已生成片段")
    store.save_job(two)
    store.recover()
    assert store.job(one["id"])["status"] == "interrupted"
    assert store.job(two["id"])["content"] == "已生成片段"


def test_template_values_are_not_recursively_expanded(tmp_path):
    from quill.legacy.template import BaseTemplatesGenerator

    renderer = BaseTemplatesGenerator(tmp_path)
    document = Document()
    paragraph = document.add_paragraph("{{变量}} 和 {{变量}}")
    renderer._replace_text_in_paragraph(paragraph, {"变量": "{{变量}}保留"})
    assert paragraph.text == "{{变量}}保留 和 {{变量}}保留"


def test_spaced_template_variables_preserve_runs_and_literal_values(tmp_path):
    from quill.legacy.template import BaseTemplatesGenerator

    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("项目：{{ 项").bold = True
    paragraph.add_run("目名称 }} / 【 负责人 】 / {{数量}}")
    BaseTemplatesGenerator(tmp_path)._fill_template(document, {
        "项目名称": "【负责人】（字面值）", "负责人": "示例乙", "数量": "45",
    })
    assert paragraph.text == "项目：【负责人】（字面值） / 示例乙 / 45"
    assert paragraph.runs[0].bold is True


def test_variable_refresh_reparses_the_saved_file_and_keeps_previous_metadata_on_failure(client):
    upload(client, "示例 · 重新提取")
    store = client.app.state.store
    template = store.template("示例 · 重新提取")
    store.update_template_variables(template["id"], [])
    refreshed = client.post("/api/refresh-template-variables", data={"template_name": template["name"]})
    assert refreshed.status_code == 200
    assert {v["name"] for v in refreshed.json()["variables"]} == {"项目名称", "负责人"}
    assert store.template(template["name"])["variables"] == refreshed.json()["variables"]
    assert store.template(template["name"])["created_at"] == template["created_at"]
    assert client.post("/api/refresh-template-variables", data={"template_name": template["name"]},
                       headers={"Authorization": ""}).status_code == 401
    (store.root / "templates" / (template["id"] + ".docx")).write_bytes(b"invalid")
    assert client.post("/api/refresh-template-variables", data={"template_name": template["name"]}).status_code == 400
    assert store.template(template["name"])["variables"] == refreshed.json()["variables"]
    assert client.post("/api/refresh-template-variables", data={"template_name": "missing"}).status_code == 404


def test_ai_template_fill_retries_field_queries_and_retains_real_sources(client):
    upload(client, "示例 · 资料取材")
    queries = []

    class Knowledge:
        async def search(self, query, collections, **options):
            queries.append((query, collections, options))
            if query in {"项目名称", "负责人"}:
                return [{"resource_id": "actual-source", "title": "项目确认单",
                         "content": "项目名称为北岸项目，负责人为示例乙。"}]
            return []

    class Model:
        async def chat(self, system, prompt):
            assert "项目确认单" in prompt
            return json.dumps({"项目名称": "北岸项目", "负责人": "示例乙"})

    client.app.state.generation.knowledge = Knowledge()
    client.app.state.llm = Model()
    filled = client.post("/api/ai-fill-template-variables", json={
        "template_name": "示例 · 资料取材", "enable_rag": True, "kb_name": "selected-collection",
    })
    assert filled.status_code == 200, filled.text
    assert filled.json()["variables"] == {"项目名称": "北岸项目", "负责人": "示例乙"}
    assert len(filled.json()["sources"]) == 1
    assert filled.json()["sources"][0]["resource_id"] == "actual-source"
    assert all(collections == ["selected-collection"] for _, collections, _ in queries)
    assert [query for query, _, _ in queries[1:]] == ["项目名称", "负责人"]
    assert all(options == {"mode": "keyword", "top_k": 3} for _, _, options in queries[1:])


async def test_template_fallback_stays_bounded_and_preserves_no_hit_failure(tmp_path):
    from quill.generation import GenerationService

    class Knowledge:
        def __init__(self):
            self.queries = []

        async def search(self, query, collections, **options):
            self.queries.append(query)
            return []

    knowledge = Knowledge()
    service = GenerationService(Store(tmp_path), knowledge, None)
    with pytest.raises(ProviderError, match="没有检索到"):
        await service.references("template", ["only-this-collection"], "", fallback_queries=[str(i) for i in range(50)])
    assert knowledge.queries == ["template"] + [str(i) for i in range(8)]


def test_word_export_contains_a_real_toc_field(tmp_path):
    from quill.legacy.document import DocumentGenerator

    path = DocumentGenerator(tmp_path).export_to_doc("# 报告\n\n## 背景\n\n正文", "fulltext")
    document = Document(path)
    fields = document._element.xpath(".//w:fldSimple")
    assert fields
    assert "背景" in "".join(document._element.xpath(".//w:fldSimple//w:t/text()"))
    assert document.core_properties.title == "报告"
    assert not document.styles['Title']._element.xpath('./w:pPr/w:pBdr')
    assert document.styles['Normal']._element.xpath("./w:rPr/w:rFonts/@w:eastAsia") == ["宋体"]
    assert "TOC" not in "\n".join(p.text for p in document.paragraphs)


def test_job_enqueued_and_worker_completes(client):
    class Model:
        async def chat(self, system, prompt):
            return "已生成的文档内容。"

    client.app.state.generation.llm = Model()
    job = client.post("/api/v1/documents/jobs", json={"topic": "工作报告", "reference": "已完成部署"})
    assert job.status_code == 202
    # The worker runs on the ASGI lifespan event loop; synchronize against its queue.
    client.portal.call(client.app.state.generation.queue.join)
    result = client.get("/api/v1/documents/jobs/" + job.json()["id"]).json()
    assert result["status"] == "completed"
    assert client.get("/download/" + result["filename"]).status_code == 200
