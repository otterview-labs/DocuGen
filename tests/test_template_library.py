import io
import json

import pytest
from docx import Document
from fastapi.testclient import TestClient

from quill.config import Settings
from quill.main import create_app
from quill.template_library import TEMPLATES


@pytest.fixture
def client(tmp_path):
    with TestClient(create_app(Settings(data_dir=tmp_path, api_key="", _env_file=None))) as local:
        yield local


@pytest.mark.parametrize("key", [template["id"] for template in TEMPLATES])
def test_starter_download_install_fill_and_word_structure(client, key):
    catalog = client.get("/api/template-library").json()["templates"]
    template = next(item for item in catalog if item["id"] == key)
    original = client.get(f"/api/template-library/{key}/download")
    assert original.status_code == 200
    assert "filename*=UTF-8''" in original.headers["content-disposition"]
    source = Document(io.BytesIO(original.content))
    assert source.paragraphs[0].text == template["name"]
    assert source.sections[0].page_width.cm == pytest.approx(21, abs=0.01)
    assert source.sections[0].page_height.cm == pytest.approx(29.7, abs=0.01)
    assert client.get("/api/templates").json()["templates"] == []
    installed = client.post(f"/api/template-library/{key}/use").json()["template"]
    assert installed["category"] == template["category"]
    assert {v["name"] for v in installed["variables"]} == {v["name"] for v in template["variables"]}
    values = {v["name"]: "已确认内容：" + v["name"] + "\n第二行" for v in installed["variables"]}
    rendered = client.post(
        "/api/generate-from-template",
        data={
            "template_name": installed["name"],
            "data": json.dumps(values),
        },
    )
    assert rendered.status_code == 200
    download = client.get("/download/" + rendered.json()["filename"])
    final = Document(io.BytesIO(download.content))
    text = "\n".join(
        [p.text for p in final.paragraphs]
        + [cell.text for table in final.tables for row in table.rows for cell in row.cells]
    )
    for value in values.values():
        assert value in text
    assert "{{" not in text
    assert final.paragraphs[0].style.name == "Title"
    assert not final.styles["Title"].element.xpath("./w:pPr/w:pBdr")
    assert final.tables[0].style.name == "Table Grid"
    assert [p.text for p in final.paragraphs if p.style.name == "Heading 1"] == [
        section["title"] for section in template["sections"]
    ]
    assert final.sections[0].header.paragraphs[0].text == template["name"]
    assert 'w:instr="PAGE"' in final.sections[0].footer._element.xml
    assert client.post(f"/api/template-library/{key}/use").json()["template"] == installed
    assert len(client.get("/api/templates").json()["templates"]) == 1


def test_install_preserves_same_named_upload_and_persists_metadata(client):
    original = client.get("/api/template-library/weekly-report/download").content
    custom = client.post(
        "/api/upload-template",
        data={
            "template_name": "工作周报",
            "category": "报告",
            "description": "自定义排版与字段",
        },
        files={"file": ("custom.docx", original)},
    )
    assert custom.status_code == 200
    installed = client.post("/api/template-library/weekly-report/use").json()["template"]
    assert installed["name"] != "工作周报"
    assert client.app.state.store.template("工作周报") == custom.json()["template"]
    assert custom.json()["template"]["category"] == "报告"
    assert custom.json()["template"]["description"] == "自定义排版与字段"
    assert len(client.get("/api/templates").json()["templates"]) == 2


def test_library_authorization_unknown_and_upload_validation(tmp_path):
    with TestClient(create_app(Settings(data_dir=tmp_path, api_key="secret", _env_file=None))) as local:
        for method, path in [
            ("GET", "/api/template-library"),
            ("GET", "/api/template-library/weekly-report/download"),
            ("POST", "/api/template-library/weekly-report/use"),
        ]:
            assert local.request(method, path).status_code == 401
        local.headers["Authorization"] = "Bearer secret"
        assert local.get("/api/template-library/missing/download").status_code == 404
        assert local.post("/api/template-library/missing/use").status_code == 404
        original = local.get("/api/template-library/weekly-report/download").content
        for fields in [{"category": "invalid"}, {"description": "长" * 501}]:
            assert (
                local.post(
                    "/api/upload-template", data=fields, files={"file": ("custom.docx", original)}
                ).status_code
                == 400
            )
